"""Perkakas penyusun dokumen Word dan PDF.

Modul ini memisahkan ISI dari BENTUK. Naskah panduan ditulis sebagai daftar
blok - judul, paragraf, tabel, kotak catatan - lalu modul ini yang mengubahnya
menjadi dokumen Word bergaya tetap, dan Word sendiri yang mengubahnya menjadi
PDF.

Pemisahan itu ada alasannya. Naskah panduan akan disunting berkali-kali sampai
batas waktu lomba; bila gaya huruf dan jarak paragraf bercampur dengan
kalimatnya, setiap penyuntingan menjadi pekerjaan tata letak. Dengan pemisahan
ini, mengubah satu kalimat berarti mengubah satu untai teks.

PDF dihasilkan lewat Microsoft Word, bukan lewat pustaka penulis PDF. Word
sudah terpasang di komputer ini, dan hasilnya identik dengan yang dilihat orang
ketika membuka berkas Word-nya - termasuk daftar isi, nomor halaman, dan
pemenggalan tabel. Pustaka PDF mana pun akan menghasilkan tata letak yang
sedikit berbeda, dan perbedaan itu baru ketahuan ketika dokumen sudah dicetak.
"""

from __future__ import annotations

import os
import subprocess
import sys
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# ---------------------------------------------------------------------------
# Palet dan ukuran
# ---------------------------------------------------------------------------
# Warna diambil dari palet aplikasi supaya dokumen dan layar terlihat berasal
# dari satu pekerjaan yang sama.
BIRU = RGBColor(0x1E, 0x5A, 0x8F)
TINTA = RGBColor(0x0F, 0x17, 0x2A)
TINTA_KEDUA = RGBColor(0x47, 0x55, 0x69)
TINTA_REDUP = RGBColor(0x94, 0xA3, 0xB8)
GARIS = "D8DEE7"
LATAR_CATATAN = "F1F5F9"
LATAR_KEPALA = "1E5A8F"
LATAR_KODE = "F8FAFC"

HURUF_ISI = "Georgia"
HURUF_JUDUL = "Segoe UI Semibold"
HURUF_KODE = "Consolas"


# ---------------------------------------------------------------------------
# Pembantu XML tingkat rendah
# ---------------------------------------------------------------------------
def _warnai_latar(elemen, warna_hex: str) -> None:
    """Beri warna latar pada sel tabel atau paragraf."""
    arsir = OxmlElement("w:shd")
    arsir.set(qn("w:val"), "clear")
    arsir.set(qn("w:color"), "auto")
    arsir.set(qn("w:fill"), warna_hex)
    elemen.append(arsir)


def _garis_sel(sel, sisi: Iterable[str], warna: str = GARIS, tebal: int = 6) -> None:
    tcPr = sel._tc.get_or_add_tcPr()
    batas = OxmlElement("w:tcBorders")
    for s in sisi:
        e = OxmlElement(f"w:{s}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(tebal))
        e.set(qn("w:color"), warna)
        batas.append(e)
    tcPr.append(batas)


def _sisipkan_medan(paragraf, instruksi: str) -> None:
    """Sisipkan medan Word (daftar isi, nomor halaman) yang dihitung Word sendiri."""
    r = paragraf.add_run()
    awal = OxmlElement("w:fldChar")
    awal.set(qn("w:fldCharType"), "begin")
    teks = OxmlElement("w:instrText")
    teks.set(qn("xml:space"), "preserve")
    teks.text = instruksi
    pisah = OxmlElement("w:fldChar")
    pisah.set(qn("w:fldCharType"), "separate")
    isi = OxmlElement("w:t")
    isi.text = "(perbarui medan: tekan Ctrl+A lalu F9)"
    akhir = OxmlElement("w:fldChar")
    akhir.set(qn("w:fldCharType"), "end")
    for e in (awal, teks, pisah, isi, akhir):
        r._r.append(e)


# ---------------------------------------------------------------------------
# Penyusun dokumen
# ---------------------------------------------------------------------------
class Panduan:
    """Penyusun dokumen dengan gaya tetap."""

    def __init__(self) -> None:
        self.doc = Document()
        self._siapkan_gaya()
        self._siapkan_halaman()

    # -- penyiapan --------------------------------------------------------
    def _siapkan_halaman(self) -> None:
        for s in self.doc.sections:
            s.top_margin = Cm(2.4)
            s.bottom_margin = Cm(2.2)
            s.left_margin = Cm(2.6)
            s.right_margin = Cm(2.2)

    def _siapkan_gaya(self) -> None:
        g = self.doc.styles["Normal"]
        g.font.name = HURUF_ISI
        g.font.size = Pt(10.5)
        g.font.color.rgb = TINTA
        pf = g.paragraph_format
        pf.space_after = Pt(7)
        pf.line_spacing = 1.28

        for nama, ukuran, atas, bawah, warna in (
            ("Heading 1", 19, 22, 8, BIRU),
            ("Heading 2", 14, 16, 6, BIRU),
            ("Heading 3", 11.5, 12, 4, TINTA),
        ):
            s = self.doc.styles[nama]
            s.font.name = HURUF_JUDUL
            s.font.size = Pt(ukuran)
            s.font.bold = True
            s.font.color.rgb = warna
            s.paragraph_format.space_before = Pt(atas)
            s.paragraph_format.space_after = Pt(bawah)
            s.paragraph_format.keep_with_next = True

    # -- blok -------------------------------------------------------------
    def sampul(self, judul: str, anak_judul: str, baris: list[tuple[str, str]]) -> None:
        for _ in range(4):
            self.doc.add_paragraph()

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(judul)
        r.font.name = HURUF_JUDUL
        r.font.size = Pt(30)
        r.font.bold = True
        r.font.color.rgb = BIRU

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        r = p.add_run(anak_judul)
        r.font.name = HURUF_JUDUL
        r.font.size = Pt(13)
        r.font.color.rgb = TINTA_KEDUA

        for _ in range(3):
            self.doc.add_paragraph()

        t = self.doc.add_table(rows=0, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for kiri, kanan in baris:
            sel = t.add_row().cells
            sel[0].width = Cm(4.6)
            sel[1].width = Cm(9.4)
            pk = sel[0].paragraphs[0]
            pk.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            rk = pk.add_run(kiri)
            rk.font.size = Pt(9.5)
            rk.font.color.rgb = TINTA_REDUP
            rk.font.name = HURUF_JUDUL
            pv = sel[1].paragraphs[0]
            rv = pv.add_run(kanan)
            rv.font.size = Pt(10)
            rv.font.color.rgb = TINTA
            rv.font.name = HURUF_JUDUL
            for s in (sel[0], sel[1]):
                s.paragraphs[0].paragraph_format.space_after = Pt(3)
        self.putus_halaman()

    def daftar_isi(self, judul: str = "Daftar Isi") -> None:
        # Judulnya sengaja TIDAK memakai gaya Heading, supaya ia tidak ikut
        # tercantum di dalam daftar isi yang dibuatnya sendiri.
        par = self.doc.add_paragraph()
        par.paragraph_format.space_after = Pt(10)
        r = par.add_run(judul)
        r.font.name = HURUF_JUDUL
        r.font.size = Pt(19)
        r.bold = True
        r.font.color.rgb = BIRU

        p = self.doc.add_paragraph()
        _sisipkan_medan(p, r'TOC \o "1-2" \h \z \u')
        self.putus_halaman()

    def h1(self, teks: str, *, tanpa_nomor: bool = False) -> None:
        if not tanpa_nomor:
            self.putus_halaman()
        self.doc.add_heading(teks, level=1)

    def h2(self, teks: str) -> None:
        self.doc.add_heading(teks, level=2)

    def h3(self, teks: str) -> None:
        self.doc.add_heading(teks, level=3)

    def p(self, teks: str) -> None:
        """Paragraf. Teks di antara ** ** dicetak tebal."""
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for bagian, tebal in _pecah_tebal(teks):
            r = par.add_run(bagian)
            r.bold = tebal

    def poin(self, butir: list[str], *, bernomor: bool = False) -> None:
        gaya = "List Number" if bernomor else "List Bullet"
        for b in butir:
            par = self.doc.add_paragraph(style=gaya)
            par.paragraph_format.space_after = Pt(3)
            par.paragraph_format.left_indent = Cm(0.7)
            for bagian, tebal in _pecah_tebal(b):
                r = par.add_run(bagian)
                r.bold = tebal

    def tabel(self, kepala: list[str], baris: list[list[str]], lebar: list[float] | None = None) -> None:
        t = self.doc.add_table(rows=1, cols=len(kepala))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = True

        for i, judul in enumerate(kepala):
            sel = t.rows[0].cells[i]
            _warnai_latar(sel._tc.get_or_add_tcPr(), LATAR_KEPALA)
            par = sel.paragraphs[0]
            par.paragraph_format.space_before = Pt(3)
            par.paragraph_format.space_after = Pt(3)
            r = par.add_run(judul)
            r.bold = True
            r.font.size = Pt(9)
            r.font.name = HURUF_JUDUL
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        for isi in baris:
            sel = t.add_row().cells
            for i, nilai in enumerate(isi[: len(kepala)]):
                par = sel[i].paragraphs[0]
                par.paragraph_format.space_before = Pt(2)
                par.paragraph_format.space_after = Pt(2)
                for bagian, tebal in _pecah_tebal(str(nilai)):
                    r = par.add_run(bagian)
                    r.bold = tebal
                    r.font.size = Pt(9)
                    r.font.name = HURUF_JUDUL
                _garis_sel(sel[i], ("top", "bottom"))

        if lebar:
            for b in t.rows:
                for i, w in enumerate(lebar[: len(kepala)]):
                    b.cells[i].width = Cm(w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def kode(self, teks: str) -> None:
        t = self.doc.add_table(rows=1, cols=1)
        sel = t.rows[0].cells[0]
        _warnai_latar(sel._tc.get_or_add_tcPr(), LATAR_KODE)
        _garis_sel(sel, ("top", "bottom", "left", "right"))
        par = sel.paragraphs[0]
        par.paragraph_format.space_before = Pt(4)
        par.paragraph_format.space_after = Pt(4)
        for i, baris in enumerate(teks.rstrip().splitlines()):
            r = par.add_run(("\n" if i else "") + baris)
            r.font.name = HURUF_KODE
            r.font.size = Pt(8.5)
            r.font.color.rgb = TINTA
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def catatan(self, judul: str, teks: str) -> None:
        t = self.doc.add_table(rows=1, cols=1)
        sel = t.rows[0].cells[0]
        _warnai_latar(sel._tc.get_or_add_tcPr(), LATAR_CATATAN)
        _garis_sel(sel, ("left",), warna="1E5A8F", tebal=18)
        par = sel.paragraphs[0]
        par.paragraph_format.space_before = Pt(5)
        par.paragraph_format.space_after = Pt(2)
        r = par.add_run(judul)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.name = HURUF_JUDUL
        r.font.color.rgb = BIRU
        par2 = sel.add_paragraph()
        par2.paragraph_format.space_after = Pt(5)
        par2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for bagian, tebal in _pecah_tebal(teks):
            r = par2.add_run(bagian)
            r.bold = tebal
            r.font.size = Pt(9.5)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def putus_halaman(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def nomor_halaman(self) -> None:
        """Pasang nomor halaman di kaki setiap halaman."""
        for s in self.doc.sections:
            p = s.footer.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run("")
            r.font.size = Pt(8.5)
            r.font.color.rgb = TINTA_REDUP
            r.font.name = HURUF_JUDUL
            _sisipkan_medan(p, "PAGE")

    def simpan(self, jalur: Path) -> Path:
        jalur.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(jalur))
        return jalur


def _pecah_tebal(teks: str) -> list[tuple[str, bool]]:
    """Pisahkan teks menurut penanda ** ** menjadi potongan biasa dan tebal."""
    bagian: list[tuple[str, bool]] = []
    sisa = teks
    while "**" in sisa:
        awal = sisa.index("**")
        if awal:
            bagian.append((sisa[:awal], False))
        sisa = sisa[awal + 2 :]
        if "**" not in sisa:
            bagian.append(("**" + sisa, False))
            return bagian
        akhir = sisa.index("**")
        bagian.append((sisa[:akhir], True))
        sisa = sisa[akhir + 2 :]
    if sisa:
        bagian.append((sisa, False))
    return bagian or [("", False)]


# ---------------------------------------------------------------------------
# Word -> PDF
# ---------------------------------------------------------------------------
def ke_pdf(docx: Path, pdf: Path | None = None) -> Path:
    """Ubah berkas Word menjadi PDF memakai Microsoft Word.

    Word dijalankan pada proses terpisah, bukan lewat pustaka ini, karena
    pustaka COM harus diinisialisasi pada utas yang sama dengan pemanggilnya -
    dan menjalankannya sebagai proses tersendiri menghindarkan seluruh
    persoalan itu sekaligus memastikan Word benar-benar tertutup sesudahnya.
    """
    pdf = pdf or docx.with_suffix(".pdf")
    skrip = (
        "import sys, pythoncom, win32com.client as w\n"
        "pythoncom.CoInitialize()\n"
        "app = w.DispatchEx('Word.Application')\n"
        "app.Visible = False\n"
        "app.DisplayAlerts = 0\n"
        "d = None\n"
        "try:\n"
        "    d = app.Documents.Open(sys.argv[1], ReadOnly=False)\n"
        "    # Daftar isi memakai medan yang baru terisi setelah dihitung ulang.\n"
        "    for t in d.TablesOfContents:\n"
        "        t.Update()\n"
        "    d.Fields.Update()\n"
        "    d.Repaginate()\n"
        "    d.SaveAs2(sys.argv[2], FileFormat=17)\n"
        "finally:\n"
        "    if d is not None:\n"
        "        d.Close(SaveChanges=-1)\n"
        "    app.Quit()\n"
        "    pythoncom.CoUninitialize()\n"
    )
    hasil = subprocess.run(
        [sys.executable, "-c", skrip, str(docx.resolve()), str(pdf.resolve())],
        capture_output=True,
        text=True,
        timeout=300,
    )
    if hasil.returncode != 0:
        raise RuntimeError(f"Word gagal membuat PDF:\n{hasil.stderr[-1500:]}")
    if not pdf.exists():
        raise RuntimeError("Word selesai tanpa galat, namun berkas PDF tidak terbentuk.")
    return pdf


__all__ = ["Panduan", "ke_pdf"]
